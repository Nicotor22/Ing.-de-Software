import io
import logging
from datetime import datetime
from xml.sax.saxutils import escape

from fastapi import HTTPException
from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from ..models import Autor, Exposicion, ExposicionObra, Obra, ObraAutor

logger = logging.getLogger(__name__)

FORMATOS_SOPORTADOS = ["pdf", "excel"]  # Requisito especial del CU-11
COLUMNAS = ["Obra", "Autor/es", "Período", "Técnica", "Estado de conservación", "Digitalizada"]
CLAVES = ["obra", "autores", "periodo", "tecnica", "estado_conservacion", "digitalizada"]
SIN_DATO = "—"


def validar_parametros(
    db: Session,
    autor_id: int | None,
    periodo: str | None,
    tecnica: str | None,
    exposicion_id: int | None = None,
) -> None:
    """Paso 2 del flujo principal: el sistema valida que los parámetros sean válidos."""
    if autor_id is not None and db.get(Autor, autor_id) is None:
        raise HTTPException(status_code=422, detail="El autor indicado no existe")
    if exposicion_id is not None and db.get(Exposicion, exposicion_id) is None:
        raise HTTPException(status_code=422, detail="La exposición indicada no existe")
    for nombre, valor in (("periodo", periodo), ("tecnica", tecnica)):
        if valor is not None and not valor.strip():
            raise HTTPException(status_code=422, detail=f"El filtro '{nombre}' no puede estar vacío")


def generar_reporte(
    db: Session,
    autor_id: int | None = None,
    periodo: str | None = None,
    tecnica: str | None = None,
    exposicion_id: int | None = None,
) -> dict:
    """Paso 3: procesa los datos según los filtros indicados."""
    validar_parametros(db, autor_id, periodo, tecnica, exposicion_id)
    periodo = periodo.strip() if periodo else None
    tecnica = tecnica.strip() if tecnica else None

    consulta = select(Obra).options(selectinload(Obra.autores).selectinload(ObraAutor.autor))
    if autor_id is not None:
        consulta = consulta.where(Obra.autores.any(ObraAutor.id_autor == autor_id))
    if periodo:
        # PROVISORIO (hasta definir qué es "periodo"): comparación exacta sin distinguir
        # mayúsculas. Con ilike("%XX%") también traería "XXI".
        consulta = consulta.where(func.lower(Obra.periodo) == periodo.lower())
    if tecnica:
        consulta = consulta.where(Obra.tecnica.ilike(f"%{tecnica}%"))
    if exposicion_id is not None:
        consulta = consulta.where(
            Obra.id.in_(
                select(ExposicionObra.id_obra).where(ExposicionObra.id_exposicion == exposicion_id)
            )
        )
    consulta = consulta.order_by(Obra.titulo)

    filas = []
    advertencias = []  # Garantía mínima: no ocultar datos inconsistentes o incompletos
    sin_periodo = 0
    sin_tecnica = 0
    for obra in db.scalars(consulta).all():
        if not obra.autores:
            advertencias.append(
                f"La obra id {obra.id} ('{obra.titulo or 'sin título'}') no tiene autores asignados"
            )
        if not obra.periodo:
            sin_periodo += 1
        if not obra.tecnica:
            sin_tecnica += 1
        filas.append(
            {
                "obra": obra.titulo or SIN_DATO,
                "autores": ", ".join(oa.autor.nombre for oa in obra.autores) or SIN_DATO,
                "periodo": obra.periodo or SIN_DATO,
                "tecnica": obra.tecnica or SIN_DATO,
                "estado_conservacion": obra.estado_conservacion or SIN_DATO,
                "digitalizada": "Sí" if obra.digitalizada else "No",
            }
        )
    if sin_periodo:
        advertencias.append(f"{sin_periodo} obra(s) sin período cargado: los resultados pueden estar incompletos")
    if sin_tecnica:
        advertencias.append(f"{sin_tecnica} obra(s) sin técnica cargada: los resultados pueden estar incompletos")

    pedidos = {"autor_id": autor_id, "periodo": periodo, "tecnica": tecnica, "exposicion_id": exposicion_id}
    return {
        "vacio": len(filas) == 0,
        # Extensión "no hay datos": se informa que el reporte está vacío
        "mensaje": "El reporte está vacío: no hay datos que cumplan el filtro" if not filas else None,
        "advertencias": advertencias,
        "filtros": {k: v for k, v in pedidos.items() if v is not None},
        "generado_en": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "total": len(filas),
        "filas": filas,
    }


def _a_matriz(reporte: dict) -> list[list]:
    return [[fila[clave] for clave in CLAVES] for fila in reporte["filas"]]


def _excel(reporte: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Reporte patrimonial"
    hoja.append(COLUMNAS)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
    for fila in _a_matriz(reporte):
        hoja.append(fila)
    for letra, ancho in zip("ABCDEF", (40, 35, 18, 25, 22, 12)):
        hoja.column_dimensions[letra].width = ancho

    # Las advertencias y los filtros viajan también en el archivo
    info = libro.create_sheet("Información")
    info.append(["Generado", reporte["generado_en"]])
    info.append(["Total de obras", reporte["total"]])
    for nombre, valor in reporte["filtros"].items():
        info.append([f"Filtro: {nombre}", valor])
    if reporte["advertencias"]:
        info.append([])
        info.append(["Advertencias"])
        info.cell(row=info.max_row, column=1).font = Font(bold=True)
        for adv in reporte["advertencias"]:
            info.append([adv])
    info.column_dimensions["A"].width = 70

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _pdf(reporte: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    estilos = getSampleStyleSheet()
    celda = estilos["BodyText"]
    celda.fontSize = 8
    celda.leading = 10

    elementos = [
        Paragraph("Reporte patrimonial", estilos["Title"]),
        Paragraph(
            f"Generado: {escape(reporte['generado_en'])} — Total de obras: {reporte['total']}",
            estilos["Normal"],
        ),
    ]
    if reporte["filtros"]:
        texto = ", ".join(f"{k}={v}" for k, v in reporte["filtros"].items())
        elementos.append(Paragraph(f"Filtros: {escape(texto)}", estilos["Normal"]))
    for adv in reporte["advertencias"]:
        elementos.append(Paragraph(f"<b>Advertencia:</b> {escape(adv)}", estilos["Normal"]))
    elementos.append(Spacer(1, 10))

    # Paragraph en cada celda: el texto largo se ajusta en vez de salirse de la hoja
    filas = [[Paragraph(f"<b>{escape(c)}</b>", celda) for c in COLUMNAS]] + [
        [Paragraph(escape(str(valor)), celda) for valor in fila] for fila in _a_matriz(reporte)
    ]
    tabla = Table(filas, repeatRows=1, colWidths=[200, 170, 70, 110, 100, 70])
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    elementos.append(tabla)

    salida = io.BytesIO()
    SimpleDocTemplate(
        salida, pagesize=landscape(A4), leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30
    ).build(elementos)
    return salida.getvalue()


def validar_formato(formato: str) -> str:
    """Extensión del paso 4: formato no disponible -> se ofrecen los soportados."""
    formato = formato.lower().strip()
    if formato not in FORMATOS_SOPORTADOS:
        raise HTTPException(
            status_code=400,
            detail={
                "mensaje": "El formato de exportación solicitado no está disponible",
                "formatos_soportados": FORMATOS_SOPORTADOS,
            },
        )
    return formato


def exportar_reporte(reporte: dict, formato: str) -> tuple[bytes, str, str]:
    """Pasos 4 y 5. Devuelve (contenido, media_type, nombre_de_archivo)."""
    formato = validar_formato(formato)
    if reporte["vacio"]:
        raise HTTPException(status_code=422, detail=reporte["mensaje"])
    try:
        if formato == "pdf":
            return _pdf(reporte), "application/pdf", "reporte_patrimonial.pdf"
        return (
            _excel(reporte),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "reporte_patrimonial.xlsx",
        )
    except Exception:
        # Extensión del paso 5: falla la generación -> se notifica el error y se puede reintentar
        logger.exception("Falló la generación del archivo de exportación")
        raise HTTPException(
            status_code=500,
            detail="No se pudo generar el archivo de exportación; podés reintentar",
        )